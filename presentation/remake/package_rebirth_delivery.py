from pathlib import Path
from pptx import Presentation
import fitz,hashlib,json,re,shutil,zipfile
R=Path(__file__).resolve().parent;P=R.parent;REPO=R.parents[1];C=R/'rebirth_system_16k';E=C/'evidence';A=C/'assets'
deck=P/'StoryBridge_智理杯_5分钟展示.pptx';pdf=P/'StoryBridge_离线展示.pdf';video=P/'StoryBridge_18秒操作演示.mp4'
shutil.copyfile(R/'preview.pdf',pdf)
src=(REPO/'backend/data/scripts/corpus_rebirth.md').read_text();target=json.loads((E/'option_d_target.json').read_text());base=(E/'baseline_first_response.md').read_text();qa=json.loads((R/'powerpoint_validation.json').read_text(encoding='utf-8-sig'));talks=json.loads((R/'talks.json').read_text())
assert len(re.findall(r'^【S\d{2}',src,re.M))==12 and '前程卡' not in src
for s in ['屏幕跳出：669','知夏，我查到487','沈曼发了条动态','你明明只有330','成绩刚出，你怎么知道我原来的分数']:assert s in src
assert len(target['scenes'])==12 and 'You saw 330' in '\n'.join(x['text'] for x in target['scenes'])
assert len(re.findall(r'^\*\*S\d{2}\s*[—-]',base,re.M))==12 and 'When you saw your score was 487' in base
sha=hashlib.sha256(deck.read_bytes()).hexdigest();prs=Presentation(deck);hidden=[i for i,s in enumerate(prs.slides,1) if s._element.get('show')=='0']
assert len(prs.slides)==12 and hidden==[11,12] and qa['pptxSha256']==sha and qa['textOverflow']==[]
assert qa['visiblePlaybackOrder']==list(range(1,11)) and qa['playback']['durationMs']==18000 and qa['playback']['positionAfter2s']>1000 and 'playbackError' not in qa
assert sum(x['seconds'] for x in talks if not x['appendix'])==300
with fitz.open(pdf) as d:assert len(d)==12 and '谢谢大家' in d[9].get_text()
dur=qa['playback']['durationMs']/1000;assert dur==18.0
report={'pptx_sha256':sha,'slides':12,'main_slides':10,'hidden_slides':[11,12],'speaker_seconds':300,'source_scenes':12,'storybridge_scenes':12,'baseline_scenes':12,'text_overflow':0,'windows_powerpoint_version':qa['applicationVersion'],'windows_video_position_after_2s_ms':qa['playback']['positionAfter2s'],'windows_native_passed':True,'mac_native_tested':False,'pdf_pages':12,'standalone_video_seconds':dur,'storybridge_internal_status':'fail','known_graph_gap':'CM10 initial propagation returned no downstream scenes'}
(R/'artifact_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
files={'StoryBridge_智理杯_5分钟展示.pptx':deck,'StoryBridge_离线展示.pdf':pdf,'StoryBridge_18秒操作演示.mp4':video,'StoryBridge_5分钟讲稿.md':P/'StoryBridge_5分钟讲稿.md','StoryBridge_展示提纲.md':P/'StoryBridge_展示提纲.md','StoryBridge_现场使用说明.md':P/'StoryBridge_现场使用说明.md','README.md':P/'README.md','素材/中文_demo_原稿.md':REPO/'backend/data/scripts/corpus_rebirth.md','素材/实验说明.md':C/'实验说明.md','素材/完整故事与改稿.md':C/'完整故事与改稿.md','素材/evidence/plans.json':E/'plans.json','素材/evidence/option_d_apply.json':E/'option_d_apply.json','素材/evidence/option_d_target.json':E/'option_d_target.json','素材/evidence/baseline_first_response.md':E/'baseline_first_response.md','素材/evidence/baseline_validation.json':E/'baseline_validation.json','素材/evidence/quote_provenance.json':E/'quote_provenance.json','素材/evidence/calls.jsonl':E/'calls.jsonl','素材/真实产品截图/system_plan.png':A/'system_plan.png','素材/真实产品截图/system_graph.png':A/'system_graph.png','素材/真实产品截图/system_s06.png':A/'system_s06.png','素材/真实产品截图/system_s10.png':A/'system_s10.png','验证记录/artifact_validation.json':R/'artifact_validation.json','验证记录/powerpoint_validation.json':R/'powerpoint_validation.json','验证记录/video_validation.txt':C/'video_validation.txt','验证记录/recording_manifest.json':C/'recording_manifest.json'}
manifest={n:{'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size} for n,p in files.items()};dest=P/'StoryBridge_展示交付包.zip'
with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
 for n,p in files.items():z.write(p,n)
 z.writestr('验证记录/文件清单.json',json.dumps(manifest,ensure_ascii=False,indent=2))
with zipfile.ZipFile(dest) as z:assert z.testzip() is None
delivery={'zip':dest.name,'zip_sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'zip_files':len(manifest)+1,'zip_bytes':dest.stat().st_size,'pptx_sha256':sha,'windows_native_passed':True,'mac_native_tested':False}
(R/'delivery_manifest.json').write_text(json.dumps(delivery,ensure_ascii=False,indent=2));print(json.dumps(delivery,ensure_ascii=False))
