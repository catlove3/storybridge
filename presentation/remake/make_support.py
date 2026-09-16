"""Build the provenance, full readable scripts, and practical offline handover."""
import hashlib,json,re,zipfile,shutil
from pathlib import Path
from collections import defaultdict
from pptx import Presentation
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A3
from PIL import Image

ROOT=Path(__file__).resolve().parent;P=ROOT/'evidence';OUT=ROOT.parent
# A user-reordered deck must not be overwritten by the historical layout's
# fixed page counts or handover text. Maintain its own order and notes instead.
if (ROOT/'deck_order.json').exists():
 from sync_current_deck import sync
 sync()
 raise SystemExit(0)
def read(name):return json.loads((P/name).read_text())
def scene(data,sid):return next(s for s in data['scenes'] if s['id']==sid)['text']
naive=(P/'naive_full.md').read_text()
bridge=read('state_v2.json');english=read('storybridge.json');source=read('state_v1.json')
naive_s02=naive.split('## S02 —')[1].split('## S03 —')[0]
assert "A photo of someone else's body with my face on it?" in naive_s02
assert 'crew team roster and seat assignment sheet' in naive
assert 'But I was in the hospital those two days.' in scene(english,'S02')
assert 'a blurry crew photo' in scene(english,'S02')
assert 'Photoshopping' in scene(english,'S05')
assert '高考' not in scene(source,'S03')
assert read('naive_response.json')['finish_reason']=='stop'
assert read('naive_protocol.json')['source_sha256']==hashlib.sha256((ROOT/'source.md').read_bytes()).hexdigest()
quotes=[('state_v2.json','S02','赛艇区域赛参赛确认函'),('state_v2.json','S03','真实队员名单和座位安排表'),('state_v2.json','S03','我记得这个位置坐的不是你。'),('state_v2.json','S05','你给我买的未来，连里面的人都不是我。'),('state_v2.json','S07','这一次，上面的名字和作品，都是我的。')]
for f,s,q in quotes:assert q in scene(read(f),s),(f,s,q)
path=next(x for x in read('propagation.json')['affected_scenes'] if x['scene_id']=='S03')
assert path['reason_path']==['CM01','E02','E03','E04','E05','E06','S03']
provenance={'main_comparison':{'stage':'Both final English scripts. Slide 4 shows Chinese excerpt translations; exact originals on slide 10.','naive_file':'naive_full.md','naive_scene':'S02','naive_exact':"Sees what? A photo of someone else's body with my face on it?",'naive_translation':'别人的身体，贴上我的脸？','storybridge_file':'storybridge.json','storybridge_scene':'S02','storybridge_exact':'But I was in the hospital those two days.','storybridge_translation':'可那两天我一直在医院。','interpretation':'Explicit face replacement is disclosed earlier in this naive response; this is an editorial observation, not a blind-rated quality result. Both workflows update the rowing clues.'},'quotes':[{'file':f,'scene':s,'exact':q} for f,s,q in quotes],'graph_path':path}
(P/'quote_provenance.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2))

calls=[json.loads(x) for x in (P/'calls.jsonl').read_text().splitlines()]
assert len([x for x in calls if x['group']=='naive_full'])==1
metrics=defaultdict(lambda:defaultdict(float))
for c in calls:
 if c['group'] not in ('storybridge','naive_full'):continue
 m=metrics[c['group']];m['calls']+=1;m['seconds']+=c['seconds'];m['prompt_tokens']+=c['response']['prompt_tokens'];m['completion_tokens']+=c['response']['completion_tokens']
report=['# 本次展示的案例与对照口径','','## 这次演示要回答的问题','','面向短剧出海编剧，演示“更换一个文化设定，如何检查被牵动的场景并审改全文”。案例是虚构七场短剧《不是我的成绩》，从高考替考改成伪造赛艇履历；人物、学校、俱乐部与赛事为虚构。','','## 当前 baseline：整份原稿交给同一个 AI','','- 输入：完整七场中文原稿、相同目标市场/英语输出要求、全部剧情保留要求，以及 StoryBridge 生成的同一选定赛艇方案。','- 系统提示只有“你是一名编剧。”；用户要求直接生成完整英语剧本，并允许附改编说明。','- 没有限制只能改哪些场景，没有关键词筛选，也没有禁止模型解释。','- 调用前保存 `naive_protocol.json`，取首次完整响应，不因结果好坏换种子或反复抽样。','- `naive_response.json` 的结束原因是 stop，全文保存在 `naive_full.md`，未经人工改稿。','','## 实际观察','','朴素全文改编也把第三场改成真实船队名单与座位表。因此，“普通 AI 漏改考场证据”这一说法已经撤回，不再出现在当前 PPT、讲稿和交付包说明中。','','PPT 第 4 页比较两边最终英语稿的中文摘译，原句可在第 10 页核对。朴素稿在第二场明确说出“别人的身体贴上我的脸”，动作描述也明确身体、姿态与光线不匹配。StoryBridge 的第二场是模糊合影与住院疑点，到了第五场才明确揭示换脸。我们将前者判断为提前释放了本应在后续揭穿场景出现的信息。这个编辑判断没有经过独立盲评，不能当作稳定质量胜率。','','朴素稿包含改编说明，并非完全没有解释。它没有返回本产品这样的可操作依赖路径与逐场前后对照；本次产品展示的亮点是创作者可以顺着影响关系审稿。没有单独验证“依赖图导致了悬念更好”的因果结论。','','## 实际图谱与检查边界','',f'第三场原文没有“高考”，实际路径为 `{" → ".join(path["reason_path"])}`。第 5 页为便于观看而压缩绘制的简图，不是完整 UI 截图；原始依据见 `propagation.json`，真实截图见 `assets/s03_impact.png`。','','初次传播覆盖 S01–S03，后续复核/修复又改写后段场景。最终内部检查状态仍为 **fail**。报告与旧结构描述存在冲突，但不能因此将所有失败一概认定为误报；完整成稿仍需人工审阅。','','准备过程还尝试过窄范围机制与关键词局部对照，已退出当前展示。旧探索记录保留在工作目录用于追溯，当前交付包仅把朴素全文改编与 StoryBridge 列为主对照。此前强提示词全文对照也更新了朋友和证据，本轮没有得出通用模型无法处理这个任务的结论。','','## 同模型、不同流程','',f'配置模型均为 `{read("naive_protocol.json")["model"]}`。StoryBridge 为现有多阶段工作流，朴素 baseline 一次调用；不是相同计算预算。温度等沿用项目配置。','', '| 组别 | 调用数 | 输入 token | 输出 token | 各调用耗时合计 / 秒 |','|---|---:|---:|---:|---:|']
for k,m in metrics.items():report.append(f'| {k} | {int(m["calls"])} | {int(m["prompt_tokens"])} | {int(m["completion_tokens"])} | {m["seconds"]:.1f} |')
report+=['','调用耗时合计不是前端等待时长，也不是重复实验平均数。不声称成本更低或已节省某个比例的人工时间。','','## 市场数据来源','','第 1 页使用 Sensor Tower 2026 年 6 月发布的 [State of Short Drama Apps 2026](https://sensortower.com/blog/state-of-short-drama-apps-2026-report)：2026 年第一季度全球短剧 App 内购收入约 7.5 亿美元。口径是全球短剧 App 内购收入估算，不是全部短剧产业收入，也不是本产品收入或市场规模。这里只用它交代赛道背景；目标用户需求仍待真实团队访谈与试用验证。','','## 视频与截图','','18 秒视频来自真实 React 前端操作，读取保存的本次真实模型结果。前半是选方案、看图谱，后半是第三场差异与结尾；模型生成等待已剪去，录屏字幕明确标注。恢复展示数据时使用 replay 任务记录，没有伪造模型文本。视频无音轨，无 BGM。原始时间段见 `recording_manifest.json`。','','## 核对与重建','','- `source.md`：完整原稿。','- `evidence/naive_prompt.md`、`naive_protocol.json`：朴素全文 baseline 的完整输入。','- `evidence/naive_full.md`、`naive_response.json`：首次真实输出。','- `evidence/state_v1.json`、`state_v2.json`、`storybridge.json`：项目改编前后中文结构与最终英语稿。','- `evidence/quote_provenance.json`：主讲引文和图路径定位。','- `evidence/calls.jsonl`：原始调用；交付包另存当前两组调用的筛选副本，不包含 API 密钥。','- `python3.12 presentation/build_ppt.py`：仅重建 PPT，不调用 API。','- `backend/.venv/bin/python presentation/remake/run_naive_baseline.py`：真实调用脚本；已存在首次结果时主动拒绝覆盖。']
(ROOT/'实验说明.md').write_text('\n'.join(report)+'\n')

story=['# 《不是我的成绩》：完整原稿与实际改稿','','模型输出原样保留。第 4 页比较两组最终英语稿，使用中文摘译；第 6 页使用 StoryBridge 中文结构稿的节选压缩。','','## 原稿','',(ROOT/'source.md').read_text(),'','## StoryBridge 中文结构改稿','']
for s in bridge['scenes']:story += [f'### {s["id"]} {s["title"]}','',s['text'],'']
story += ['## 朴素全文改编：完整英语结果','',naive,'','## StoryBridge：完整英语结果','']
for s in english['scenes']:story += [f'### {s["id"]} {s["title"]}','',s['text'],'']
(ROOT/'完整故事与改稿.md').write_text('\n'.join(story)+'\n')

deck=OUT/'StoryBridge_智理杯_5分钟展示.pptx';prs=Presentation(deck)
assert len(prs.slides)==12
hidden=[i for i,s in enumerate(prs.slides,1) if s._element.get('show')=='0'];assert hidden==[9,10,11,12]
bounds=[]
for i,s in enumerate(prs.slides,1):
 for sh in s.shapes:
  if sh.left < -100 or sh.top < -100 or sh.left+sh.width>prs.slide_width+500 or sh.top+sh.height>prs.slide_height+500:bounds.append((i,sh.name))
assert not bounds,bounds
talks=json.loads((ROOT/'talks.json').read_text());assert sum(x['seconds'] for x in talks)==300
all_text='\n'.join(sh.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
assert '谢谢大家' in all_text
assert '只改直接出现' not in all_text
with zipfile.ZipFile(deck) as z:
 assert z.testzip() is None
 media=[x for x in z.namelist() if x.startswith('ppt/media/') and x.endswith('.mp4')];assert len(media)==1
 assert z.read(media[0])==(OUT/'StoryBridge_18秒操作演示.mp4').read_bytes()
 assert not [x for x in z.namelist() if x.startswith('ppt/media/') and x.lower().endswith(('.mp3','.wav','.m4a','.aac'))]
 for name in z.namelist():
  if name.endswith('.rels'):
   from xml.etree import ElementTree as ET
   for rel in ET.fromstring(z.read(name)):
    if rel.get('Type','').endswith(('/video','/audio','/media')):assert rel.get('TargetMode')!='External'
validation={'slides':12,'main_slides':8,'hidden_slides':hidden,'out_of_bounds':bounds,'speaker_notes_seconds':300,'embedded_video_matches_standalone':True,'quotes_verified':len(quotes)+3,'graph_path_verified':True,'naive_baseline_uses_first_complete_response':True,'script_chinese_characters':sum(len(re.findall(r'[\u4e00-\u9fff]',x['text'])) for x in talks if not x['appendix']),'pptx_sha256':hashlib.sha256(deck.read_bytes()).hexdigest()}
(ROOT/'artifact_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2))

# Place the unchanged generated poster on an A3 PDF page; no raster editing.
poster=OUT/'StoryBridge_产品海报.png'
if poster.exists():
 c=canvas.Canvas(str(OUT/'StoryBridge_产品海报_A3.pdf'),pagesize=A3)
 pw,ph=A3;iw,ih=Image.open(poster).size;factor=min(pw/iw,ph/ih)
 c.drawImage(str(poster),(pw-iw*factor)/2,(ph-ih*factor)/2,width=iw*factor,height=ih*factor)
 c.setTitle('StoryBridge 产品海报');c.showPage();c.save()

handover='''# 现场使用说明

## 直接用哪份

打开 `StoryBridge_智理杯_5分钟展示.pptx`。共 8 页主讲、4 页隐藏备答。第 7 页点击播放 18 秒视频，第 8 页“谢谢大家”结束。各页已标页码，讲稿建议总时长 5 分钟，包含视频。

`StoryBridge_5分钟讲稿.md` 包含逐页讲法与时间段；PPT 备注中也有相同内容。用“谁用—原稿—任务—对照—依赖图—改后—操作—结尾”记住主线，现场对着 PPT 讲。控制第 4 页讲解时间，避免在 baseline 上讲成技术报告。

## 现场脱网也能讲

1. 正常方案：本地 PPT，视频已经嵌入，无须联网或打开产品服务。
2. 视频不能播放：直接打开独立 `StoryBridge_18秒操作演示.mp4`；或者放映时输入 `12` 后回车，讲真实静态对照。结束后输入 `8` 回车回到致谢。
3. PPT 格式异常：打开 `StoryBridge_离线展示.pdf`。前 8 页是主讲，后 4 页为备答；第 7 页是视频封面，需另开 MP4 或跳到第 12 页。

上台前把 PPT、PDF 和 MP4 放在电脑本地同一目录，不要依赖聊天软件的在线预览。PDF 中没有可播放视频。无需插入背景音乐。

## 设备检查

- 已在 Windows 原生 PowerPoint 16.0 导出检查；具体幻灯片、文本溢出与实际视频播放记录见验证文件。
- 同源浏览器预览也检查了文本边界；它不替代 Office 原生检查。
- 当前没有可用的 Mac 实机，**Mac 版 PowerPoint / Keynote 尚未验证**。拿到 Mac 后打开最终 PPT，重点看第 1、4、6 页中文换行和第 7 页视频；如有问题用 PDF + 独立 MP4。
- 中文使用微软雅黑。PDF 已固定排版，能作为缺字体或软件差异时的备用。
- PPT 为本地原生排版，没有模板平台的 AI 水印；视频没有音轨或 BGM。生成海报没有可见 AI 水印，原图元数据原样保留。

## 主讲与答辩

一人主讲；另一人协助翻页、视频与计时；其余成员准备依赖图、baseline 与产品边界的答辩。技术页全部放在谢谢大家后面，不挤占 5 分钟。主讲人按实际语速完整排练两次，穿着整洁，对观众说话，不逐字读稿。

- 第 9 页：公平对照的完整输入与同模型口径。
- 第 10 页：两边最终英语原句及我们的编辑判断。
- 第 11 页：真实路径、内部检查 fail 与还需人工审稿的边界。
- 第 12 页：视频故障时的静态截图。

不要说“普通 AI 漏改了考场材料”或“我们已经全面超过大模型”。本轮朴素全文改编也改对了赛艇材料；主讲观察到的是它更早说破换脸。依赖图提供检查路径，尚无大样本盲评证明稳定质量提升。

## 素材与海报

海报 PNG 使用内置 image_gen 生成；A3 PDF 保持原图，仅放到 A3 页面。提示词与生成方式见 `poster_prompt.md`。源码、截图与原始模型输出保留在工作目录，`实验说明.md` 解释每个比较的口径。手机扫码体验本次暂缓，未在 PPT 上放尚未上线的二维码。
'''
(OUT/'StoryBridge_现场使用说明.md').write_text(handover)
readme='''# StoryBridge · 当前展示稿

直接打开 **StoryBridge_智理杯_5分钟展示.pptx**。本目录只保留一个当前 PPT；旧展示稿已清理。

8 页主讲 + 4 页隐藏备答。顺序是：产品与用户 → 完整原故事 → 美国翻拍任务 → 整篇交给 AI 的真实对照 → 依赖图 → 改写结果 → 18 秒实际操作 → 谢谢大家。

第 7 页内嵌视频，第 8 页结束，第 9–12 页为答辩备用。PDF、独立 MP4 和第 12 页截图用于脱网备用。Mac 实机尚未验证。

Baseline 是完整原稿、相同要求、相同方案交给同一模型一次改编；本轮它也改对了赛艇线索。PPT 展示的是两组实际结果的悬念次序与产品可查看的审改流程，不再使用关键词局部基线。

讲稿和现场使用说明都在本目录。完整包还包括原生 PDF、海报、完整故事、实验说明与可核对证据。

重建：`python3.12 presentation/build_ppt.py`，只读取素材，不调用 API。原始运行和验证文件位于 `remake/`。产品核心代码未修改；扫码体验暂缓。
'''
(OUT/'README.md').write_text(readme)
print(json.dumps(validation,ensure_ascii=False))
