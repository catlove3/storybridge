"""Update navigation metadata in the user's saved PPTX without rebuilding slides."""
from __future__ import annotations
import copy,hashlib,io,json,os,re,shutil,tempfile,zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent
DECK=OUT/'StoryBridge_智理杯_5分钟展示.pptx'
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','ep':'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}

def title_of(slide):
 texts=[sh for sh in slide.shapes if sh.has_text_frame and sh.text.strip()]
 if any(sh.text=='谢谢大家' for sh in texts):return '谢谢大家'
 if any(sh.shape_type==16 for sh in slide.shapes):return '18 秒看一次实际改编操作'
 heads=[sh for sh in texts if .7<float(sh.top)/914400<1.6]
 return max(heads,key=lambda s:(s.text_frame.paragraphs[0].font.size or 0)).text if heads else texts[0].text

def role_of(title):
 for role,phrase in [('thanks','谢谢大家'),('video','18 秒看'),('protocol','把所有改编信息'),('quotes','第二场的英语原句'),('graph','Dependency graph'),('adapted','观众能跟着线索'),('limits','成稿仍要由人审'),('comparison','第二场就说破了'),('fallback','视频备用'),('original','她的名字'),('task','制作方'),('intro','StoryBridge')]:
  if phrase in title:return role
 return 'other'

def set_text(shape,value):
 """Keep the existing paragraph/run style for short navigation text."""
 ts=shape._element.xpath('.//a:t')
 if not ts:raise ValueError('Text shape has no text run')
 ts[0].text=value
 for e in ts[1:]:e.text=''

def sync():
 original=DECK.read_bytes();before_hash=hashlib.sha256(original).hexdigest()
 prs=Presentation(io.BytesIO(original))
 prior=Presentation(ROOT/'presentation.pptx')
 order=[s.slide_id for s in prs.slides]
 new_page={s.slide_id:i for i,s in enumerate(prs.slides,1)}
 remap={i:new_page[s.slide_id] for i,s in enumerate(prior.slides,1) if s.slide_id in new_page}
 end=next(i for i,s in enumerate(prs.slides,1) if title_of(s)=='谢谢大家')
 roles={role_of(title_of(s)):i for i,s in enumerate(prs.slides,1)}
 video=roles['video'];fallback=roles['fallback']
 def refs(value):
  return re.sub(r'第\s*(\d+)\s*页',lambda m:f'第 {remap.get(int(m[1]),int(m[1]))} 页',value)

 base_times={'intro':35,'original':35,'task':35,'protocol':15,'quotes':25,'graph':40,'adapted':35,'video':40,'limits':15,'thanks':25,'comparison':45,'fallback':0,'other':20}
 weights=[base_times[role_of(title_of(s))] for s in list(prs.slides)[:end]]
 raw=[w*300/sum(weights) for w in weights];times=[int(v) for v in raw]
 for i in sorted(range(end),key=lambda i:raw[i]-times[i],reverse=True)[:300-sum(times)]:times[i]+=1
 promoted={
  'protocol':'我们把完整七场原稿、目标市场、全部改编要求，连同一个赛艇方案，都交给同一个模型，让它一次写完英语剧本。没有限制它只能改哪些场景，也允许它附改编说明。下面看两边实际写出了什么。',
  'quotes':'左边是朴素全文改编第二场的原句：女儿已经说出“别人的身体贴上我的脸”。右边 StoryBridge 的同一场只有“那两天我在医院”，先留下疑点，到第五场再拿原始名单确认。这次两边都改对了赛艇线索，我们关注的是反转何时说破。',
  'limits':'图能解释哪些场景互相牵动，但本轮内部检查仍未通过。我们保留了原始报告，完整成稿还需要人工审阅；接下来也需要用更多类型的剧本进行盲评。',
 }
 patches={};changes=[];rows=[];talks=[]
 for i,s in enumerate(prs.slides,1):
  role=role_of(title_of(s));hidden=i>end;dirty=False
  if (s._element.get('show')=='0')!=hidden:
   s._element.set('show','0' if hidden else '1');dirty=True
  for sh in s.shapes:
   if not sh.has_text_frame:continue
   old=sh.text;new=old
   if re.fullmatch(r'\d{2}\s*/\s*\d{2}|备答\s*\d+',old.strip()):
    new=f'备答 {i-end}' if hidden else f'{i:02d} / {end:02d}'
   elif not hidden and old.startswith('备答：'):
    new=old.removeprefix('备答：')
   elif '第' in old and '页' in old:
    new=refs(old)
   if old!=new:
    set_text(sh,new);dirty=True;changes.append({'page':i,'shape_id':sh.shape_id,'before':old,'after':new})
  if dirty:patches[str(s.part.partname).lstrip('/')]=etree.tostring(s._element,encoding='UTF-8',xml_declaration=True,standalone=True)
  old_notes=s.notes_slide.notes_text_frame.text
  m=re.match(r'建议用时：\s*(\d+)\s*秒\s*',old_notes)
  prose=refs(old_notes[m.end():] if m else old_notes)
  if not hidden and role in promoted and m and int(m[1])==0:prose=promoted[role]
  if role=='video' and not hidden:
   prose='下面看十八秒实际操作。请留意：选赛艇方案、看第三场为什么受影响、对照改前改后的文字。【点击视频，18 秒。】朋友的经历和她手里的名单一起改了，女儿靠画作申请的结尾仍然保留。这是实际产品读取真实结果的回放，省略了生成等待。'
  if role=='fallback':
   prose=f'截图来自实际项目的 S03 前后对照。放映时按 {fallback} 回车跳到本页，然后按 {end} 回车回到谢谢大家。PDF 第 {video} 页为视频封面，可打开独立 MP4，或直接讲本页截图。'
  seconds=0 if hidden else times[i-1]
  new_notes=f'建议用时：{seconds} 秒\n\n{prose}'
  if new_notes!=old_notes:
   s.notes_slide.notes_text_frame.text=new_notes
   patches[str(s.notes_slide.part.partname).lstrip('/')]=etree.tostring(s.notes_slide._element,encoding='UTF-8',xml_declaration=True,standalone=True)
  title=title_of(s)
  rows.append({'page':i,'slide_id':s.slide_id,'title':title,'role':role,'hidden':hidden,'seconds':seconds,'footer':f'备答 {i-end}' if hidden else f'{i:02d} / {end:02d}'})
  talks.append({'page':i,'seconds':seconds,'text':prose,'title':title,'appendix':hidden})

 with zipfile.ZipFile(io.BytesIO(original)) as z:
  app=etree.fromstring(z.read('docProps/app.xml'))
  for tag,value in [('Slides',len(rows)),('HiddenSlides',len(rows)-end),('Notes',len(rows)),('MMClips',1)]:
   el=app.find('ep:'+tag,NS)
   if el is not None:el.text=str(value)
  patches['docProps/app.xml']=etree.tostring(app,encoding='UTF-8',xml_declaration=True,standalone=True)
  tmp=DECK.with_suffix('.syncing.pptx')
  with zipfile.ZipFile(tmp,'w') as target:
   for item in z.infolist():target.writestr(copy.copy(item),patches.get(item.filename,z.read(item.filename)))
  with zipfile.ZipFile(tmp) as target:
   assert target.testzip() is None
   # Media, layouts, relationships and all other untouched parts remain byte-identical.
   assert all(target.read(n)==z.read(n) for n in z.namelist() if n not in patches)
   assert target.read('ppt/presentation.xml')==z.read('ppt/presentation.xml')
  assert [s.slide_id for s in Presentation(tmp).slides]==order
  if hashlib.sha256(DECK.read_bytes()).hexdigest()!=before_hash:
   tmp.unlink();raise RuntimeError('PPT changed while syncing; rerun against the latest saved file.')
  backup=Path(tempfile.gettempdir())/f'storybridge-before-navigation-{before_hash[:12]}.pptx'
  if not backup.exists():backup.write_bytes(original)
  os.replace(tmp,DECK)
 shutil.copyfile(DECK,ROOT/'presentation.pptx')
 final_hash=hashlib.sha256(DECK.read_bytes()).hexdigest()
 manifest={'source':'Current user-edited PPTX; order and visual content preserved.','pptx_sha256':final_hash,'total_slides':len(rows),'main_slides':end,'hidden_slides':[r['page'] for r in rows if r['hidden']],'thanks_slide':end,'video_slide':video,'role_pages':roles,'speaker_notes_seconds':sum(times),'slides':rows}
 (ROOT/'deck_order.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
 (ROOT/'talks.json').write_text(json.dumps(talks,ensure_ascii=False,indent=2))
 (ROOT/'navigation_sync.json').write_text(json.dumps({'before_sha256':before_hash,'after_sha256':final_hash,'order_preserved':True,'slide_ids':order,'unchanged_parts_byte_identical':True,'changes':changes},ensure_ascii=False,indent=2))

 elapsed=0;md=['# StoryBridge · 5 分钟讲稿','',f'{end} 页主讲、{len(rows)-end} 页隐藏备答；按当前 PPT 顺序。第 {video} 页含 18 秒视频，建议总用时 300 秒，包含视频。','']
 for talk in talks:
  stop=elapsed+talk['seconds'];md += [f'## {talk["page"]}. {talk["title"]}','']
  if not talk['appendix']:md += [f'{elapsed//60}:{elapsed%60:02d}–{stop//60}:{stop%60:02d}','']
  else:md += ['隐藏备答；不计入主讲时间。','']
  md += [talk['text'],''];elapsed=stop
 (OUT/'StoryBridge_5分钟讲稿.md').write_text('\n'.join(md))
 outline=['# StoryBridge · 当前页面顺序','',f'前 {end} 页可见，后 {len(rows)-end} 页隐藏；第 {end} 页谢谢大家结束。','', '| 页码 | 内容 | 放映 | 建议用时 |','|---:|---|---|---:|']
 outline += [f'| {r["page"]} | {r["title"]} | {"隐藏备答" if r["hidden"] else "主讲可见"} | {r["seconds"]} 秒 |' for r in rows]
 (OUT/'StoryBridge_展示提纲.md').write_text('\n'.join(outline)+'\n')

 # Keep evidence and prose unchanged; only repair their references to slide numbers.
 for name in ['实验说明.md','完整故事与改稿.md']:
  f=ROOT/name;f.write_text(refs(f.read_text()))
 p=ROOT/'evidence/quote_provenance.json';q=json.loads(p.read_text())
 q['main_comparison']['stage']=f'Both final English scripts. Slide {roles["comparison"]} shows Chinese excerpt translations; exact originals on slide {roles["quotes"]}.'
 p.write_text(json.dumps(q,ensure_ascii=False,indent=2))
 validation=json.loads((ROOT/'artifact_validation.json').read_text())
 validation.update({'slides':len(rows),'main_slides':end,'hidden_slides':manifest['hidden_slides'],'speaker_notes_seconds':300,'pptx_sha256':final_hash,'user_slide_order_preserved':True,'untouched_parts_byte_identical':True,'script_chinese_characters':sum(len(re.findall(r'[\u4e00-\u9fff]',t['text'])) for t in talks if not t['appendix'])})
 (ROOT/'artifact_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2))
 write_handover(manifest)
 print(json.dumps({'main_slides':end,'hidden_slides':manifest['hidden_slides'],'graph_slide':roles['graph'],'video_slide':video,'thanks_slide':end,'speaker_notes_seconds':300,'order_preserved':True},ensure_ascii=False))

def write_handover(m):
 r=m['role_pages'];end=m['thanks_slide'];video=m['video_slide'];hidden=m['hidden_slides'];fallback=r['fallback']
 main_order=' → '.join(s['title'] for s in m['slides'] if not s['hidden'])
 readme=f'''# StoryBridge · 当前展示稿

直接打开 **StoryBridge_智理杯_5分钟展示.pptx**。当前文件按用户调整后的页面顺序同步，前 {end} 页可见，后 {len(hidden)} 页隐藏。

依赖图在第 {r['graph']} 页，18 秒视频在第 {video} 页，“谢谢大家”在第 {end} 页。第 {r['comparison']} 页为对照摘译备答，第 {fallback} 页为视频备用截图。

讲稿、页码、跳页提示、PDF 与交付包均按这个顺序同步。建议主讲 5 分钟，含视频；Mac 实机未验证。

维护入口：`python3.12 presentation/build_ppt.py`。它读取当前 PPT，仅同步页码、可见性、备注与讲稿，保留页面顺序和现有内容，不调用 API。`remake/build_deck.py` 为历史排版生成器，不用于覆盖当前手工编辑稿。
'''
 (OUT/'README.md').write_text(readme)
 guide=f'''# 现场使用说明

## 当前页数和播放顺序

{end} 页主讲 + {len(hidden)} 页隐藏备答。第 {r['protocol']} 页介绍 baseline，第 {r['quotes']} 页给出英语原句，第 {r['graph']} 页展示依赖图，第 {video} 页播放视频，第 {r['limits']} 页说明当前边界，第 {end} 页谢谢大家结束。第 {r['comparison']}、{fallback} 页隐藏，答辩时按需打开。

讲稿按当前页面顺序安排，建议总时长 300 秒，含 18 秒视频。PPT 备注、讲稿 Markdown 和提纲的时间段一致；仍需按主讲人的实际语速排练。

## 现场备用

1. 正常使用本地 PPT，视频已嵌入，无需联网。
2. 第 {video} 页视频不能播放时，打开独立 MP4，或输入 `{fallback}` 回车跳到静态截图；结束后输入 `{end}` 回车回到谢谢大家。
3. PPT 格式异常时打开 `StoryBridge_离线展示.pdf`。共 {m['total_slides']} 页，前 {end} 页为主讲，后 {len(hidden)} 页为备答。PDF 第 {video} 页是视频封面，不能直接播放视频。
4. 原始英语引用在第 {r['quotes']} 页；需要中文大字对照时，输入 `{r['comparison']}` 回车打开隐藏备答。

把 PPT、PDF、MP4 放在电脑本地。视频无音轨，也没有 BGM。不要依赖聊天软件的在线预览。

## 检查与分工

以最终文件的 Windows 原生 PowerPoint 检查记录为准；预览图片和 PDF 来自同一文件的原生导出。Mac 版 PowerPoint / Keynote 尚未实测，拿到 Mac 后检查中文换行和第 {video} 页视频，异常时使用 PDF + 独立 MP4。

一人主讲，其他成员协助翻页、计时和答辩。按提纲排练，展示时看着观众与 PPT 讲，控制 baseline 和边界页的用时。

## 案例口径

Baseline 使用完整原稿、全部相同改编要求和同一模型一次改编。本轮它也改对了赛艇线索，比较的是实际输出的悬念顺序与项目提供的检查路径，不能称其漏改考场证据。内部检查 fail 的说明在第 {r['limits']} 页，完整成稿仍需人工审阅。

原始输入与响应、依赖图依据见 `实验说明.md` 和 evidence。海报使用内置 image_gen，提示词见 `poster_prompt.md`；扫码体验本次暂缓。
'''
 (OUT/'StoryBridge_现场使用说明.md').write_text(guide)

if __name__=='__main__':sync()
