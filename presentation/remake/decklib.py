"""Native editable PowerPoint elements plus a same-source browser preview."""
from pathlib import Path
import html,json
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.util import Inches,Pt
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parent
FONT='Microsoft YaHei'
NAVY,CREAM,WHITE='#152D3C','#FAF9F5','#FFFFFF'
INK,MUTED,RED,GOLD,GREEN='#182B37','#68757B','#C64B3C','#B79456','#267868'
PALE,LINE='#E4EEE8','#DEDACF'

def rgb(v):return RGBColor.from_string(v.lstrip('#'))

class Deck:
 def __init__(self,main_slide_count=8):
  self.prs=Presentation();self.prs.slide_width=Inches(13.333333);self.prs.slide_height=Inches(7.5)
  self.pages=[];self.talks=[];self.main_slide_count=main_slide_count
 def slide(self,title,chapter,source='',dark=False,appendix=False):
  self.current=self.prs.slides.add_slide(self.prs.slide_layouts[6])
  self.pages.append({'title':title,'items':[],'dark':dark,'appendix':appendix})
  if appendix:self.current._element.set('show','0')
  self.rect(0,0,13.333333,7.5,NAVY if dark else CREAM)
  self.text('StoryBridge   ·   '+chapter,.62,.34,12,.27,11,'#D9C39C' if dark else MUTED,True)
  if title:self.text(title,.62,.91,12.0,.96,31,WHITE if dark else INK,True)
  self.text(source or '智理杯 · 跨文化故事改编',.62,7.02,11.2,.32,10,'#AFC3C9' if dark else MUTED)
  n=f'备答 {len(self.pages)-self.main_slide_count}' if appendix else f'{len(self.pages):02d} / {self.main_slide_count:02d}'
  self.text(n,12.08,7.02,.77,.28,10,'#AFC3C9' if dark else MUTED)
  return self.current
 def rect(self,x,y,w,h,fill,line=None,round_=False):
  sh=self.current.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
  sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
  if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.7)
  else:sh.line.fill.background()
  if round_:sh.adjustments[0]=.05
  self.pages[-1]['items'].append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,line=line,round=round_))
  return sh
 def text(self,value,x,y,w,h,size=22,color=INK,bold=False,align='left'):
  sh=self.current.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
  tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
  for i,line in enumerate(value.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
   p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
   p.space_before=p.space_after=Pt(0);p.line_spacing=1.16
   p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
   # Explicit East Asian face avoids Office substituting a different CJK font.
   ea=OxmlElement('a:ea');ea.set('typeface',FONT);p._p.get_or_add_pPr().get_or_add_defRPr().append(ea)
  self.pages[-1]['items'].append(dict(kind='text',text=value,x=x,y=y,w=w,h=h,size=size,color=color,bold=bold,align=align))
  return sh
 def picture(self,path,x,y,w,h):
  path=Path(path);im=Image.open(path);ratio=min(w/im.width,h/im.height);fw,fh=im.width*ratio,im.height*ratio
  px,py=x+(w-fw)/2,y+(h-fh)/2
  self.current.shapes.add_picture(str(path),Inches(px),Inches(py),width=Inches(fw),height=Inches(fh))
  self.pages[-1]['items'].append(dict(kind='image',path=path.as_uri(),x=px,y=py,w=fw,h=fh))
 def arrow(self,x1,y1,x2,y2,color=GOLD):
  sh=self.current.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(2)
  end=OxmlElement('a:tailEnd');end.set('type','triangle');sh.line._get_or_add_ln().append(end)
  self.pages[-1]['items'].append(dict(kind='arrow',x1=x1,y1=y1,x2=x2,y2=y2,color=color))
 def notes(self,seconds,content):
  self.current.notes_slide.notes_text_frame.text=f'建议用时：{seconds} 秒\n\n{content}'
  self.talks.append({'seconds':seconds,'text':content,'title':self.pages[-1]['title'],'appendix':self.pages[-1]['appendix']})
 def movie(self,path,poster,x,y,w,h):
  sh=self.current.shapes.add_movie(str(path),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4')
  self.pages[-1]['items'].append(dict(kind='video',path=Path(path).as_uri(),poster=Path(poster).as_uri(),x=x,y=y,w=w,h=h))
  return sh
 def save(self,path):
  self.prs.save(path);self.html(ROOT/'preview.html')
  (ROOT/'layout.json').write_text(json.dumps(self.pages,ensure_ascii=False,indent=2))
  (ROOT/'talks.json').write_text(json.dumps(self.talks,ensure_ascii=False,indent=2))
 def html(self,path):
  out=['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>StoryBridge 视频对比版</title><style>',
   '@font-face{font-family:"Microsoft YaHei";src:url("file:///mnt/c/Windows/Fonts/msyh.ttc")}@font-face{font-family:"Microsoft YaHei";src:url("file:///mnt/c/Windows/Fonts/msyhbd.ttc");font-weight:700}',
   '*{box-sizing:border-box}body{margin:0;background:#aaa;font-family:"Microsoft YaHei"}.slide{width:1280px;height:720px;position:relative;overflow:hidden;margin:0 auto 24px;break-after:page}.item{position:absolute}.text{white-space:pre-wrap;overflow-wrap:break-word;line-height:1.16}@media print{@page{size:13.333333in 7.5in;margin:0}.slide{margin:0}video{display:none}}','</style><body>']
  for page in self.pages:
   out.append('<section class="slide">')
   for i,item in enumerate(page['items']):
    if item['kind']=='arrow':
     out.append(f'<svg class="item" style="left:0;top:0;width:1280px;height:720px" viewBox="0 0 1280 720"><defs><marker id="a{i}" markerWidth="7" markerHeight="7" refX="6" refY="3.5" orient="auto"><path d="M0,0 L7,3.5 L0,7 Z" fill="{item["color"]}"/></marker></defs><line x1="{item["x1"]*96}" y1="{item["y1"]*96}" x2="{item["x2"]*96}" y2="{item["y2"]*96}" stroke="{item["color"]}" stroke-width="2.5" marker-end="url(#a{i})"/></svg>');continue
    css=f'left:{item["x"]*96}px;top:{item["y"]*96}px;width:{item["w"]*96}px;height:{item["h"]*96}px;'
    if item['kind']=='rect':
     css+=f'background:{item["fill"]};'
     if item['line']:css+=f'border:1px solid {item["line"]};'
     if item['round']:css+='border-radius:8px;'
     out.append(f'<div class="item" style="{css}"></div>')
    elif item['kind']=='text':
     css+=f'font-size:{item["size"]*96/72}px;color:{item["color"]};font-weight:{700 if item["bold"] else 400};text-align:{item.get("align","left")};'
     out.append(f'<div class="item text" style="{css}">{html.escape(item["text"])}</div>')
    elif item['kind']=='video':
     out.append(f'<video class="item" style="{css}" src="{item["path"]}" poster="{item["poster"]}" controls preload="metadata"></video>')
    else:out.append(f'<img class="item" style="{css}" src="{item["path"]}">')
   out.append('</section>')
  out.append('</body></html>');Path(path).write_text('\n'.join(out))
